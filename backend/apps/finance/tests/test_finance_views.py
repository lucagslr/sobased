"""Categories, summary, budget, advances, recurring expenses, exports and the
dashboard money widgets."""

import io
import zipfile
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.accounts.tests.factories import UserFactory
from apps.contacts.models import Contact, ProjectContact
from apps.finance import services
from apps.finance.models import BudgetLine, Category, RecurringExpense, Transaction
from apps.finance.tasks import generate_recurring_expenses
from apps.projects.tests.factories import grant

from .conftest import TODAY, client_for, days, make_transaction, png_upload

pytestmark = pytest.mark.django_db


# --- Categories -----------------------------------------------------------------------


def test_every_workspace_gets_the_default_categories(tree, treasurer_api):
    admin = UserFactory(username="wsadmin")
    client = client_for(admin)

    created = client.post("/api/workspaces/", {"name": "Perso", "color": "#BFDBFE"})

    names = list(
        Category.objects.filter(workspace_id=created.data["id"])
        .order_by("position")
        .values_list("name", flat=True)
    )
    assert names[:3] == ["Studio", "Logiciel", "Matériel"] and "Autre" in names
    assert len(names) == 15


def test_categories_are_read_by_members_and_written_by_admins(tree, categories):
    admin = UserFactory(username="admin")
    grant(admin, tree.workspace, "admin")
    editor = UserFactory(username="editor")
    grant(editor, tree.workspace, "editor")
    outsider = UserFactory(username="outsider")

    listed = client_for(editor).get(f"/api/categories/?workspace={tree.workspace.pk}")
    refused = client_for(editor).post(
        "/api/categories/", {"workspace": tree.workspace.pk, "name": "Presse"}
    )
    created = client_for(admin).post(
        "/api/categories/", {"workspace": tree.workspace.pk, "name": "Presse"}
    )
    duplicate = client_for(admin).post(
        "/api/categories/", {"workspace": tree.workspace.pk, "name": "presse"}
    )

    assert len(listed.data) == 15
    assert refused.status_code == 403 and created.status_code == 201
    assert duplicate.status_code == 400
    assert client_for(outsider).get("/api/categories/").data == []


def test_deleting_a_used_category_reassigns_everything(tree, categories):
    admin = UserFactory(username="admin")
    grant(admin, tree.workspace, "admin", can_view_finance=True, can_edit_finance=True)
    studio, other, autre = (
        categories["Studio"],
        categories["Location"],
        categories["Autre"],
    )
    item = make_transaction(tree["A"], studio)
    BudgetLine.objects.create(
        project=tree["A"], category=studio, kind="expense", amount=100
    )
    BudgetLine.objects.create(
        project=tree["A"], category=other, kind="expense", amount=50
    )
    recurring = RecurringExpense.objects.create(
        project=tree["A"], label="x", amount=10, category=studio, start_date=TODAY
    )
    client = client_for(admin)

    explicit = client.delete(f"/api/categories/{studio.pk}/?replace_with={other.pk}")

    assert explicit.status_code == 204
    item.refresh_from_db()
    recurring.refresh_from_db()
    assert item.category == other and recurring.category == other
    # Budget lines merged into the replacement.
    line = BudgetLine.objects.get(project=tree["A"], category=other, kind="expense")
    assert line.amount == Decimal("150")
    # Without replace_with, "Autre" takes over.
    make_transaction(tree["A"], other)
    assert client.delete(f"/api/categories/{other.pk}/").status_code == 204
    assert Transaction.objects.filter(category=autre).count() == 2
    # "Autre" itself, while used and without any replacement, cannot go.
    for category in Category.objects.filter(workspace=tree.workspace).exclude(
        pk=autre.pk
    ):
        category.delete()
    assert client.delete(f"/api/categories/{autre.pk}/").status_code == 400


def test_deleting_a_workspace_takes_its_bookkeeping_with_it(tree, categories):
    """RESTRICT, not PROTECT, on the category keys: found the hard way."""
    owner = tree.owner
    make_transaction(tree["A"], categories["Studio"])
    BudgetLine.objects.create(
        project=tree["A"], category=categories["Studio"], kind="expense", amount=1
    )
    RecurringExpense.objects.create(
        project=tree["A"],
        label="x",
        amount=1,
        category=categories["Studio"],
        start_date=TODAY,
    )

    response = client_for(owner).delete(f"/api/workspaces/{tree.workspace.pk}/")

    assert response.status_code == 204
    assert not Transaction.objects.exists()
    assert not Category.objects.filter(workspace_id=tree.workspace.pk).exists()


# --- Summary --------------------------------------------------------------------------


def test_summary_totals_by_category_project_and_month(tree, categories, treasurer_api):
    make_transaction(
        tree["A"], categories["Studio"], amount=100, date=date(2026, 1, 10)
    )
    make_transaction(
        tree["A1"], categories["Studio"], amount=50, date=date(2026, 1, 20)
    )
    make_transaction(
        tree["B"],
        categories["Subvention"],
        amount=1000,
        kind="income",
        date=date(2026, 2, 1),
    )
    make_transaction(
        tree["B"], categories["Logiciel"], amount=30, payment_status="to_pay"
    )
    make_transaction(tree["Z"], tree.other_workspace.categories.first(), amount=999)

    data = treasurer_api.get("/api/finance/summary/").data

    assert data["expense"] == "180.00" and data["income"] == "1000.00"
    assert data["balance"] == "820.00" and data["count"] == 4
    assert data["needs_receipt"] == 3  # every expense here lacks a receipt
    assert data["to_pay"] == {"total": "30.00", "count": 1}
    assert {c["name"]: c["expense"] for c in data["by_category"]} == {
        "Logiciel": "30.00",
        "Studio": "150.00",
        "Subvention": "0.00",
    }
    assert {p["project"]: p["balance"] for p in data["by_project"]} == {
        tree["A"].pk: "-100.00",
        tree["A1"].pk: "-50.00",
        tree["B"].pk: "970.00",
    }
    january = next(m for m in data["by_month"] if m["month"] == "2026-01-01")
    assert january["expense"] == "150.00"
    # The same filters as the list: a project and its sub-projects.
    scoped = treasurer_api.get(
        f"/api/finance/summary/?project={tree['A'].pk}&include_descendants=true"
    ).data
    assert scoped["expense"] == "150.00" and scoped["income"] == "0.00"


# --- Budget ---------------------------------------------------------------------------


def test_budget_planned_versus_actual_own_and_with_children(
    tree, categories, treasurer_api
):
    studio, grant_cat = categories["Studio"], categories["Subvention"]
    client = treasurer_api

    def post(project, category, kind, amount):
        return client.post(
            "/api/budget-lines/",
            {
                "project": project.pk,
                "category": category.pk,
                "kind": kind,
                "amount": amount,
            },
            format="json",
        )

    assert post(tree["A"], studio, "expense", "500").status_code == 201
    assert post(tree["A1"], studio, "expense", "200").status_code == 201
    assert post(tree["A"], grant_cat, "income", "1000").status_code == 201
    # Posting the same line again updates it (the form is a grid).
    again = post(tree["A"], studio, "expense", "600")
    assert again.status_code == 201
    assert BudgetLine.objects.filter(project=tree["A"], category=studio).count() == 1
    make_transaction(tree["A"], studio, amount=120)
    make_transaction(tree["A1x"], studio, amount=80)
    make_transaction(tree["B"], studio, amount=999)  # another branch

    data = client.get(f"/api/finance/budget/?project={tree['A'].pk}").data

    lines = {(line["name"], line["kind"]): line for line in data["lines"]}
    studio_line = lines[("Studio", "expense")]
    assert studio_line["planned"] == "600.00" and studio_line["actual"] == "120.00"
    assert studio_line["planned_with_children"] == "800.00"
    assert studio_line["actual_with_children"] == "200.00"
    assert lines[("Subvention", "income")]["planned"] == "1000.00"
    assert data["totals"]["own"]["planned_expense"] == "600.00"
    assert data["totals"]["with_children"]["actual_expense"] == "200.00"
    assert data["totals"]["own"]["planned_income"] == "1000.00"


def test_budget_ignores_sub_projects_i_may_not_see_money_of(tree, categories):
    treasurer = UserFactory(username="t")
    grant(treasurer, tree["A"], "editor", can_view_finance=True, can_edit_finance=True)
    make_transaction(tree["A"], categories["Studio"], amount=10)
    make_transaction(tree["A1"], categories["Studio"], amount=20)
    # A guest-only sub-project? Rights are inherited downward, so A1 counts.
    # What must NOT count: a project the treasurer has no money rights on.
    other = UserFactory(username="other")
    grant(other, tree["B"], "editor", can_view_finance=True)
    make_transaction(tree["B"], categories["Studio"], amount=500)

    data = (
        client_for(treasurer).get(f"/api/finance/budget/?project={tree['A'].pk}").data
    )
    root = client_for(treasurer).get(f"/api/finance/budget/?project={tree['R'].pk}")

    assert data["totals"]["with_children"]["actual_expense"] == "30.00"
    assert root.status_code == 404  # no finance rights on R itself


# --- Advances: who owes what to whom --------------------------------------------------


def test_advances_grouped_by_payer_and_root_project(tree, categories, treasurer):
    helder = UserFactory(username="helder")
    grant(helder, tree["A1"], "editor")
    ana = Contact.objects.create(
        workspace=tree.workspace, first_name="Ana", last_name="D"
    )
    ProjectContact.objects.create(project=tree["A"], contact=ana)
    make_transaction(
        tree["A1x"],
        categories["Transport"],
        amount=40,
        to_reimburse=True,
        paid_by_user=helder,
    )
    make_transaction(
        tree["B"],
        categories["Transport"],
        amount=60,
        to_reimburse=True,
        paid_by_user=helder,
    )
    make_transaction(
        tree["A"],
        categories["Studio"],
        amount=200,
        to_reimburse=True,
        paid_by_contact=ana,
    )
    make_transaction(
        tree["A"],
        categories["Studio"],
        amount=999,
        to_reimburse=True,
        paid_by_user=helder,
        reimbursed_on=TODAY,
    )
    client = client_for(treasurer)

    data = client.get("/api/finance/advances/").data

    by_name = {person["payer_name"]: person for person in data}
    assert by_name["Ana D"]["total"] == "200.00"
    helder_entry = by_name[helder.display_name]
    assert helder_entry["total"] == "100.00"
    assert [(p["project_name"], p["total"]) for p in helder_entry["projects"]] == [
        ("R", "100.00")
    ]
    ids = helder_entry["projects"][0]["transactions"]
    # « Marquer remboursé » for the whole balance.
    settled = client.post(
        "/api/finance/advances/", {"transactions": ids}, format="json"
    )
    assert settled.status_code == 200 and settled.data["updated"] == 2
    assert client.get("/api/finance/advances/").data[0]["payer_name"] == "Ana D"


def test_settling_advances_needs_edit_rights(
    tree, categories, finance_viewer, treasurer
):
    item = make_transaction(
        tree["A"], categories["Studio"], to_reimburse=True, paid_by_user=treasurer
    )

    response = client_for(finance_viewer).post(
        "/api/finance/advances/", {"transactions": [item.pk]}, format="json"
    )

    assert response.status_code == 200 and response.data["updated"] == 0
    item.refresh_from_db()
    assert item.reimbursed_on is None


# --- Recurring expenses ---------------------------------------------------------------


def test_recurring_expense_crud_and_next_due(tree, categories, treasurer_api):
    created = treasurer_api.post(
        "/api/recurring-expenses/",
        {
            "project": tree["A"].pk,
            "label": "Studio One",
            "amount": "19.90",
            "category": categories["Logiciel"].pk,
            "frequency": "monthly",
            "day": 31,
            "start_date": "2026-01-01",
        },
        format="json",
    )

    assert created.status_code == 201
    assert created.data["next_due"] is not None
    bad_day = treasurer_api.post(
        "/api/recurring-expenses/",
        {
            "project": tree["A"].pk,
            "label": "x",
            "amount": "1",
            "category": categories["Logiciel"].pk,
            "day": 32,
            "start_date": "2026-01-01",
        },
        format="json",
    )
    assert bad_day.status_code == 400
    listed = treasurer_api.get(f"/api/recurring-expenses/?project={tree['A'].pk}")
    assert [e["label"] for e in listed.data] == ["Studio One"]


def test_nightly_generation_is_idempotent_and_respects_the_day(tree, categories):
    expense = RecurringExpense.objects.create(
        project=tree["A"],
        label="Studio One",
        amount=Decimal("19.90"),
        category=categories["Logiciel"],
        frequency="monthly",
        day=31,
        start_date=date(2026, 1, 15),
    )
    yearly = RecurringExpense.objects.create(
        project=tree["A"],
        label="Assurance",
        amount=Decimal("300"),
        category=categories["Frais admin"],
        frequency="yearly",
        day=1,
        month=3,
        start_date=date(2026, 1, 1),
        end_date=date(2026, 12, 31),
    )

    # January 15th: day 31 not reached yet, nothing.
    assert services.generate_due(date(2026, 1, 15)) == 0
    # January 31st: the monthly one, "À payer", dated on its day.
    assert services.generate_due(date(2026, 1, 31)) == 1
    assert services.generate_due(date(2026, 1, 31)) == 0  # idempotent
    january = Transaction.objects.get(recurring_expense=expense, period_key="2026-01")
    assert january.payment_status == "to_pay" and january.date == date(2026, 1, 31)
    assert january.amount == Decimal("19.90") and january.label == "Studio One"
    # February has 28 days: day 31 becomes the 28th.
    assert services.generate_due(date(2026, 2, 28)) == 1
    assert Transaction.objects.get(period_key="2026-02").date == date(2026, 2, 28)
    # March 1st: the yearly one too.
    assert services.generate_due(date(2026, 3, 1)) == 1
    assert Transaction.objects.get(recurring_expense=yearly).period_key == "2026"
    # Next year: the yearly one ended, the monthly one goes on.
    assert services.generate_due(date(2027, 3, 31)) == 1
    assert not Transaction.objects.filter(recurring_expense=yearly, period_key="2027")
    # Inactive: nothing.
    expense.is_active = False
    expense.save()
    assert services.generate_due(date(2027, 4, 30)) == 0
    assert generate_recurring_expenses() == 0


# --- Exports --------------------------------------------------------------------------


def test_exports_cover_the_filtered_transactions_only(tree, categories, treasurer_api):
    make_transaction(tree["A"], categories["Studio"], label="Studio A", amount=100)
    with_receipt = make_transaction(
        tree["A1"], categories["Transport"], label="Train", amount=20
    )
    treasurer_api.put(
        f"/api/transactions/{with_receipt.pk}/receipt/",
        {"receipt": png_upload("cff.png")},
        format="multipart",
    )
    make_transaction(tree["B"], categories["Cachet"], label="Cachet B", amount=300)
    make_transaction(
        tree["Z"], tree.other_workspace.categories.first(), label="Z secret", amount=1
    )
    query = f"?project={tree['A'].pk}&include_descendants=true"

    xlsx = treasurer_api.get(f"/api/finance/export.xlsx{query}")
    pdf = treasurer_api.get(f"/api/finance/export.pdf{query}")
    archive = treasurer_api.get(f"/api/finance/export-receipts.zip{query}")

    assert xlsx.status_code == 200
    assert xlsx["Content-Disposition"].startswith("attachment; filename=")
    book = load_workbook(io.BytesIO(xlsx.content))
    assert book.sheetnames == ["Transactions", "Par catégorie", "Par projet", "Export"]
    rows = list(book["Transactions"].iter_rows(min_row=2, values_only=True))
    assert sorted(row[4] for row in rows) == ["Studio A", "Train"]
    assert pdf.status_code == 200 and pdf.content[:5] == b"%PDF-"
    assert archive.status_code == 200
    with zipfile.ZipFile(io.BytesIO(archive.content)) as zipped:
        names = zipped.namelist()
        index = zipped.read("index.csv").decode("utf-8-sig")
    assert len([n for n in names if n.startswith("justificatifs/")]) == 1
    assert "Studio A" in index and "À justifier" in index and "Cachet B" not in index


def test_exports_are_invisible_without_finance_rights(tree, categories, blind_editor):
    make_transaction(tree["A"], categories["Studio"])
    client = client_for(blind_editor)

    xlsx = client.get("/api/finance/export.xlsx")

    assert xlsx.status_code == 200
    rows = list(
        load_workbook(io.BytesIO(xlsx.content))["Transactions"].iter_rows(min_row=2)
    )
    assert rows == []


# --- Dashboard widgets and project overview -------------------------------------------


def test_money_widgets(tree, categories, treasurer_api, blind_editor):
    make_transaction(
        tree["A"], categories["Logiciel"], label="due", payment_status="to_pay"
    )
    make_transaction(
        tree["A"],
        categories["Logiciel"],
        label="next month",
        payment_status="to_pay",
        date=days(45),
    )
    make_transaction(tree["B"], categories["Studio"], label="no receipt")

    widgets = treasurer_api.get("/api/dashboard/summary/").data["widgets"]
    blind = client_for(blind_editor).get("/api/dashboard/summary/").data["widgets"]

    assert widgets["expenses_to_pay"]["available"] is True
    assert [t["label"] for t in widgets["expenses_to_pay"]["items"]] == ["due"]
    assert widgets["missing_receipts"]["count"] == 3
    assert blind["expenses_to_pay"] == {"available": False, "count": 0}
    assert blind["missing_receipts"] == {"available": False, "count": 0}


def test_project_overview_budget_only_with_finance_rights(
    tree, categories, treasurer_api, blind_editor
):
    BudgetLine.objects.create(
        project=tree["A"], category=categories["Studio"], kind="expense", amount=500
    )
    make_transaction(
        tree["A1"], categories["Studio"], amount=75, payment_status="to_pay"
    )

    mine = treasurer_api.get(f"/api/projects/{tree['A'].pk}/overview/").data["budget"]
    blind = client_for(blind_editor).get(f"/api/projects/{tree['A'].pk}/overview/").data

    assert mine["own"]["planned_expense"] == "500.00"
    assert mine["with_children"]["actual_expense"] == "75.00"
    assert mine["needs_receipt"] == 1 and mine["to_pay"] == "75.00"
    assert blind["budget"] is None


def test_cards_carry_the_expense_budget_only_with_finance_rights(
    tree, categories, treasurer_api, blind_editor
):
    BudgetLine.objects.create(
        project=tree["A1"], category=categories["Studio"], kind="expense", amount=300
    )
    make_transaction(tree["A1x"], categories["Studio"], amount=120)

    card = treasurer_api.get("/api/projects/cards/").data[0]
    blind = client_for(blind_editor).get("/api/projects/cards/").data[0]

    assert card["budget"] == {"planned": "300.00", "spent": "120.00"}
    columns = card["past"] + card["current"] + card["upcoming"]
    entry = next(e for e in columns if e["name"] == "A")
    assert entry["budget"] == {"planned": "300.00", "spent": "120.00"}
    assert blind["budget"] is None
