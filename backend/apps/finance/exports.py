"""Exports (SPEC §13): Excel (openpyxl), PDF (WeasyPrint), ZIP of receipts.

All three take a queryset that the view already limited to the user's rights
and to the requested period: an export never contains a franc the user could
not see on screen. Amounts are formatted the Swiss way (1'234.50 CHF) and
dates as 12.10.2026 (SPECIFICATIONS §9).
"""

from __future__ import annotations

import csv
import io
import zipfile
from datetime import date
from decimal import Decimal

from django.template.loader import render_to_string

from . import services
from .models import Kind, Transaction

KIND_LABELS = dict(Kind.choices)
STATUS_LABELS = {
    "to_pay": "À payer",
    "needs_receipt": "À justifier",
    "to_reimburse": "À rembourser",
    "ok": "OK",
}


def chf(amount: Decimal | int | None) -> str:
    """1'234.50 CHF"""
    value = Decimal(amount or 0).quantize(Decimal("0.01"))
    sign = "-" if value < 0 else ""
    whole, cents = f"{abs(value):.2f}".split(".")
    grouped = f"{int(whole):,}".replace(",", "'")
    return f"{sign}{grouped}.{cents} CHF"


def swiss_date(day: date | None) -> str:
    return day.strftime("%d.%m.%Y") if day else ""


def period_of(request) -> dict:
    """The period the export covers, for its header."""
    return {
        "date_after": request.query_params.get("date_after") or "",
        "date_before": request.query_params.get("date_before") or "",
    }


def _period_label(period: dict) -> str:
    def fmt(value: str) -> str:
        return swiss_date(date.fromisoformat(value)) if value else ""

    start, end = fmt(period.get("date_after", "")), fmt(period.get("date_before", ""))
    if start and end:
        return f"du {start} au {end}"
    if start:
        return f"dès le {start}"
    if end:
        return f"jusqu'au {end}"
    return "toutes dates"


def _payer(item: Transaction) -> str:
    if item.paid_by_user_id:
        return item.paid_by_user.display_name
    if item.paid_by_contact_id:
        return item.paid_by_contact.display_name
    return ""


def _rows(transactions):
    return transactions.select_related(
        "project", "category", "contact", "paid_by_user", "paid_by_contact"
    ).order_by("date", "id")


# --- Excel ----------------------------------------------------------------------------


def build_xlsx(transactions, period: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    data = services.summary(transactions)
    book = Workbook()
    bold = Font(bold=True)

    sheet = book.active
    sheet.title = "Transactions"
    headers = [
        "Date",
        "Projet",
        "Nature",
        "Catégorie",
        "Libellé",
        "Fournisseur / contact",
        "Montant CHF",
        "Statut",
        "Payé par",
        "Remboursé le",
        "Justificatif",
    ]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = bold
    for item in _rows(transactions):
        sheet.append(
            [
                item.date,
                item.project.name,
                KIND_LABELS[item.kind],
                item.category.name,
                item.label,
                item.vendor or (item.contact.display_name if item.contact_id else ""),
                float(item.amount) if item.kind == Kind.INCOME else -float(item.amount),
                STATUS_LABELS[item.display_status],
                _payer(item),
                item.reimbursed_on,
                "oui" if item.receipt else ("manquant" if item.needs_receipt else ""),
            ]
        )
    for column in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 18
    for row in sheet.iter_rows(min_row=2):
        row[0].number_format = "DD.MM.YYYY"
        row[6].number_format = "#,##0.00"
        row[9].number_format = "DD.MM.YYYY"

    def synthesis(title: str, key: str, label: str, entries: list[dict]):
        page = book.create_sheet(title)
        page.append([label, "Dépenses CHF", "Recettes CHF", "Solde CHF"])
        for cell in page[1]:
            cell.font = bold
        for entry in entries:
            page.append(
                [
                    entry[key],
                    float(entry["expense"]),
                    float(entry["income"]),
                    float(entry["balance"]),
                ]
            )
        page.append(
            [
                "Total",
                float(data["expense"]),
                float(data["income"]),
                float(data["balance"]),
            ]
        )
        for cell in page[page.max_row]:
            cell.font = bold
        page.column_dimensions["A"].width = 32
        for letter in "BCD":
            page.column_dimensions[letter].width = 16
        for row in page.iter_rows(min_row=2):
            for cell in row[1:]:
                cell.number_format = "#,##0.00"

    synthesis("Par catégorie", "name", "Catégorie", data["by_category"])
    names = {
        row["project_id"]: row["project__name"]
        for row in transactions.values("project_id", "project__name")
    }
    synthesis(
        "Par projet",
        "name",
        "Projet",
        [{**entry, "name": names[entry["project"]]} for entry in data["by_project"]],
    )
    info = book.create_sheet("Export")
    info.append(["Période", _period_label(period)])
    info.append(["Transactions", data["count"]])
    info.append(["Dépenses sans justificatif", data["needs_receipt"]])
    info.column_dimensions["A"].width = 28

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- PDF ------------------------------------------------------------------------------


def build_pdf(transactions, period: dict, user) -> bytes:
    from weasyprint import HTML

    data = services.summary(transactions)
    names = {
        row["project_id"]: row["project__name"]
        for row in transactions.values("project_id", "project__name")
    }

    def money_rows(entries, name_of):
        return [
            {
                "name": name_of(entry),
                "expense": chf(entry["expense"]),
                "income": chf(entry["income"]),
                "balance": chf(entry["balance"]),
            }
            for entry in entries
        ]

    rows = [
        {
            "date": swiss_date(item.date),
            "project": item.project.name,
            "category": item.category.name,
            "label": item.label,
            "vendor": item.vendor,
            "status": STATUS_LABELS[item.display_status],
            "alert": item.display_status != "ok",
            "amount": ("−" if item.kind == Kind.EXPENSE else "+") + chf(item.amount),
            "needs_receipt": item.needs_receipt,
        }
        for item in _rows(transactions)
    ]
    # Django templates cannot call functions: everything is formatted here.
    html = render_to_string(
        "finance/report.html",
        {
            "period": _period_label(period),
            "generated_on": swiss_date(date.today()),
            "author": user.display_name,
            "totals": {
                "income": chf(data["income"]),
                "expense": chf(data["expense"]),
                "balance": chf(data["balance"]),
            },
            "by_category": money_rows(data["by_category"], lambda e: e["name"]),
            "by_project": money_rows(data["by_project"], lambda e: names[e["project"]]),
            "rows": rows,
            "missing": [row for row in rows if row["needs_receipt"]],
        },
    )
    return HTML(string=html).write_pdf()


# --- ZIP of receipts ------------------------------------------------------------------


def build_receipts_zip(transactions) -> bytes:
    """Every receipt, named after the transaction, plus an index CSV that
    also lists the expenses still missing one."""
    buffer = io.BytesIO()
    index = io.StringIO()
    writer = csv.writer(index, delimiter=";")
    writer.writerow(["Date", "Projet", "Libellé", "Montant CHF", "Fichier", "Statut"])
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for item in _rows(transactions):
            if item.kind != Kind.EXPENSE:
                continue
            name = ""
            if item.receipt:
                extension = item.receipt.name.rsplit(".", 1)[-1]
                name = f"{item.date.isoformat()}_{item.pk}.{extension}"
                with item.receipt.open("rb") as handle:
                    archive.writestr(f"justificatifs/{name}", handle.read())
            writer.writerow(
                [
                    swiss_date(item.date),
                    item.project.name,
                    item.label,
                    f"{item.amount:.2f}",
                    name,
                    "ok" if name else "À justifier",
                ]
            )
        archive.writestr("index.csv", "﻿" + index.getvalue())  # BOM: Excel reads UTF-8
    return buffer.getvalue()
